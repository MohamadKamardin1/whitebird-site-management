import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.accounts.factories import UserFactory
from apps.accounts.models import RoleCode
from apps.site_management.hr_services import build_cleaner_workbook_template, import_cleaner_workbook, preview_cleaner_workbook
from apps.site_management.models import Cleaner, CleanerStatus, TraineeProgram


def test_role_codes_include_hr_and_store_manager() -> None:
    assert RoleCode.HR.value == "hr"
    assert RoleCode.STORE_MANAGER.value == "store_manager"


@pytest.mark.django_db
def test_hr_template_contains_server_controlled_configuration_and_validations(site) -> None:
    content = build_cleaner_workbook_template(onboarding_status=CleanerStatus.TRAINEE, site=site)
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Onboarding configuration", "Cleaners", "ReferenceData"]
    assert workbook["Onboarding configuration"]["B2"].value == "trainee"
    assert workbook["Onboarding configuration"]["B3"].value == site.name
    assert workbook["ReferenceData"].sheet_state == "hidden"
    assert len(workbook["Cleaners"].data_validations.dataValidation) >= 3
    assert workbook["Cleaners"]["A1"].value == "first_name"
    assert "site_code" not in [cell.value for cell in workbook["Cleaners"][1]]


def _workbook_with_row(row: list[object], site) -> SimpleUploadedFile:
    wb = load_workbook(io.BytesIO(build_cleaner_workbook_template(onboarding_status=CleanerStatus.TRAINEE, site=site)))
    wb["Cleaners"].append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile("cleaners.xlsx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


_VALID_ROW = [
    "Fatma", "Ali", "nida", "HRNA0001", "female", "1991-05-05", "Stone Town",
    "+255700000001", "Omar", "Brother", "+255700000002", "2026-08-16", "",
]


@pytest.mark.django_db
def test_hr_preview_returns_row_level_errors_without_committing(site) -> None:
    row = list(_VALID_ROW)
    row[2] = "bad_id_type"
    result = preview_cleaner_workbook(_workbook_with_row(row, site))
    assert result["total_rows"] == 1
    assert result["accepted_rows"] == 0
    assert result["rejected_rows"] == 1
    assert "id_type is not an allowed value" in result["errors"][0]["errors"]


@pytest.mark.django_db
def test_hr_preview_accepts_valid_row(site) -> None:
    result = preview_cleaner_workbook(_workbook_with_row(list(_VALID_ROW), site))
    assert result["total_rows"] == 1
    assert result["accepted_rows"] == 1
    assert result["rejected_rows"] == 0


@pytest.mark.django_db(transaction=True)
def test_hr_import_creates_configured_trainee_and_site_assignment(site, admin_user) -> None:
    result = import_cleaner_workbook(
        uploaded_file=_workbook_with_row(list(_VALID_ROW), site), actor=admin_user,
        onboarding_status=CleanerStatus.TRAINEE, site=site,
    )
    assert result["committed"] is True
    cleaner = Cleaner.objects.get(id_number="HRNA0001")
    assert cleaner.status == CleanerStatus.TRAINEE
    assert TraineeProgram.objects.filter(cleaner=cleaner, site=site).exists()
    assert cleaner.site_assignments.filter(site=site, status="draft").exists()


@pytest.mark.django_db(transaction=True)
def test_hr_import_creates_configured_active_cleaner_and_assignment(site, admin_user) -> None:
    row = list(_VALID_ROW)
    row[3] = "HRNA0002"
    result = import_cleaner_workbook(
        uploaded_file=_workbook_with_row(row, site), actor=admin_user,
        onboarding_status=CleanerStatus.ACTIVE, site=site,
    )
    assert result["committed"] is True
    cleaner = Cleaner.objects.get(id_number="HRNA0002")
    assert cleaner.status == CleanerStatus.ACTIVE
    assert cleaner.site_assignments.filter(site=site, status="active").exists()
    assert not TraineeProgram.objects.filter(cleaner=cleaner).exists()


@pytest.mark.django_db(transaction=True)
def test_hr_import_rejects_duplicates_and_bad_dates(site) -> None:
    wb = load_workbook(io.BytesIO(build_cleaner_workbook_template(onboarding_status=CleanerStatus.TRAINEE, site=site)))
    wb["Cleaners"].append(_VALID_ROW)
    wb["Cleaners"].append(_VALID_ROW)
    buffer = io.BytesIO(); wb.save(buffer)
    preview = preview_cleaner_workbook(SimpleUploadedFile("dup.xlsx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
    assert preview["accepted_rows"] == 1
    assert preview["rejected_rows"] == 1
    bad = list(_VALID_ROW); bad[5] = "not-a-date"
    imported = import_cleaner_workbook(
        uploaded_file=_workbook_with_row(bad, site), actor=UserFactory(role=RoleCode.SYSTEM_ADMIN),
        onboarding_status=CleanerStatus.TRAINEE, site=site,
    )
    assert imported["committed"] is False
    assert not Cleaner.objects.filter(id_number="HRNA0001").exists()


@pytest.mark.django_db
def test_hr_template_rejects_inactive_destination_site(site) -> None:
    site.is_active = False
    site.save(update_fields=["is_active"])
    with pytest.raises(Exception):
        build_cleaner_workbook_template(onboarding_status=CleanerStatus.TRAINEE, site=site)
