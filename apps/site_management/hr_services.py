"""HR cleaner workbook generation, validation, and controlled import services."""

from __future__ import annotations

import hashlib
import io
from datetime import date, timedelta
from typing import Any

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from apps.accounts.models import User
from apps.core.models import AuditLog
from apps.core.services import record_audit

from .assignment_services import assign_cleaner_to_site, end_assignment
from .models import (
    Cleaner,
    CleanerAssignmentStatus,
    CleanerAssignmentType,
    CleanerSiteAssignment,
    CleanerStatus,
    Gender,
    IdType,
    Site,
    TraineeProgram,
    TraineeProgramStatus,
    WorkMode,
)
from .trainee_services import start_trainee_program, transfer_trainee_program


HEADERS = [
    "first_name", "last_name", "id_type", "id_number", "gender", "birth_date",
    "living_location", "phone_number", "near_person_name", "near_person_relationship",
    "near_person_phone", "registration_date", "notes",
]
REQUIRED = {"first_name", "last_name", "id_type", "id_number", "gender", "birth_date", "registration_date"}
_ALLOWED_ONBOARDING_STATUSES = {CleanerStatus.TRAINEE, CleanerStatus.ACTIVE}


def _configuration(*, onboarding_status: CleanerStatus | str, site: Site) -> CleanerStatus:
    status = CleanerStatus(onboarding_status)
    if status not in _ALLOWED_ONBOARDING_STATUSES:
        raise ValidationError("HR onboarding status must be trainee or active.")
    if not site.is_active:
        raise ValidationError("Choose an active destination site.")
    return status


def build_cleaner_workbook_template(*, onboarding_status: CleanerStatus | str, site: Site) -> bytes:
    """Build a workbook whose placement/status configuration is visible but server-owned."""
    status = _configuration(onboarding_status=onboarding_status, site=site)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cleaners"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:M201"
    for column, width in {
        "A": 18, "B": 18, "C": 20, "D": 22, "E": 14, "F": 14, "G": 24,
        "H": 18, "I": 20, "J": 20, "K": 18, "L": 16, "M": 30,
    }.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True, color="FFFFFF")
        cell.fill = cell.fill.copy(fill_type="solid", fgColor="0F7667")
    table = Table(displayName="CleanerOnboardingRows", ref="A1:M201")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)

    reference = workbook.create_sheet("ReferenceData")
    reference.sheet_state = "hidden"
    lists = {"A": ("id_type", [item.value for item in IdType]), "B": ("gender", [item.value for item in Gender])}
    for column, (label, values) in lists.items():
        reference[f"{column}1"] = label
        for index, value in enumerate(values, 2):
            reference[f"{column}{index}"] = value

    configuration = workbook.create_sheet("Onboarding configuration", 0)
    configuration.append(["White Bird HR Cleaner Onboarding Workbook"])
    configuration.append(["Selected onboarding status", status.value])
    configuration.append(["Selected assigned site", site.name])
    configuration.append(["Selected site code", site.code])
    configuration.append(["Status and site are selected by HR before download and are enforced at preview/import."])
    configuration.append(["Complete the Cleaners sheet. Dates must be YYYY-MM-DD. Do not add site or status columns."])
    configuration.append(["Use the dropdowns for id_type and gender. Upload is preview-first: rejected rows block import."])
    for row in configuration.iter_rows():
        row[0].alignment = row[0].alignment.copy(wrap_text=True)
    configuration.column_dimensions["A"].width = 105
    configuration.column_dimensions["B"].width = 32

    validations = [
        DataValidation(type="list", formula1="=ReferenceData!$A$2:$A$4", allow_blank=False),
        DataValidation(type="list", formula1="=ReferenceData!$B$2:$B$5", allow_blank=False),
        DataValidation(type="date", operator="between", formula1="DATE(1940,1,1)", formula2="TODAY()", allow_blank=False),
    ]
    for validation in validations:
        sheet.add_data_validation(validation)
    validations[0].add("C2:C201")
    validations[1].add("E2:E201")
    validations[2].add("F2:F201")
    validations[2].add("L2:L201")
    sheet.conditional_formatting.add(
        "A2:M201", FormulaRule(formula=["COUNTIF($D$2:$D$201,$D2)>1"], fill=sheet["A2"].fill.copy(fill_type="solid", fgColor="FCE4D6"))
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _normalise(value: Any) -> str:
    return str(value or "").strip()


def _parse_date(value: Any, field: str, row_number: int) -> date:
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(_normalise(value))
    except ValueError as exc:
        raise ValidationError(f"Row {row_number}: {field} must be a valid YYYY-MM-DD date.") from exc


def _rows(uploaded_file: Any) -> tuple[list[dict[str, Any]], str]:
    raw = uploaded_file.read()
    file_hash = hashlib.sha256(raw).hexdigest()
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if "Cleaners" not in workbook.sheetnames:
        raise ValidationError("Workbook must contain a Cleaners sheet.")
    sheet = workbook["Cleaners"]
    headers = [_normalise(value).lower() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    if headers != HEADERS:
        raise ValidationError("Cleaners sheet headers do not match the current White Bird template.")
    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value not in (None, "") for value in values):
            continue
        row = dict(zip(HEADERS, values, strict=False))
        row["_row_number"] = row_number
        rows.append(row)
    return rows, file_hash


def _validate_row(row: dict[str, Any], *, seen_ids: set[tuple[str, str]]) -> tuple[dict[str, Any] | None, list[str]]:
    number = int(row["_row_number"])
    errors: list[str] = []
    cleaned: dict[str, Any] = {key: _normalise(value) for key, value in row.items() if not key.startswith("_")}
    for field in REQUIRED:
        if not cleaned.get(field):
            errors.append(f"{field} is required")
    if cleaned.get("id_type") and cleaned["id_type"] not in {item.value for item in IdType}:
        errors.append("id_type is not an allowed value")
    if cleaned.get("gender") and cleaned["gender"] not in {item.value for item in Gender}:
        errors.append("gender is not an allowed value")
    if cleaned.get("id_type") and cleaned.get("id_number"):
        key = (cleaned["id_type"], cleaned["id_number"].upper())
        if key in seen_ids:
            errors.append("duplicate identity appears more than once in this workbook")
        seen_ids.add(key)
    for field in ("birth_date", "registration_date"):
        if cleaned.get(field):
            try:
                cleaned[field] = _parse_date(row[field], field, number)
            except ValidationError as exc:
                errors.append(str(exc))
    return (cleaned if not errors else None), errors


def preview_cleaner_workbook(uploaded_file: Any) -> dict[str, Any]:
    rows, file_hash = _rows(uploaded_file)
    accepted: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for row in rows:
        cleaned, row_errors = _validate_row(row, seen_ids=seen)
        if row_errors:
            errors.append({"row": row["_row_number"], "errors": row_errors})
        else:
            accepted.append({"row": row["_row_number"], "data": cleaned})
    return {"file_hash": file_hash, "total_rows": len(rows), "accepted_rows": len(accepted), "rejected_rows": len(errors), "accepted": accepted, "errors": errors}


def import_cleaner_workbook(*, uploaded_file: Any, actor: User, onboarding_status: CleanerStatus | str, site: Site) -> dict[str, Any]:
    """Import one HR-configured cohort; spreadsheet cells cannot alter its site or status."""
    status = _configuration(onboarding_status=onboarding_status, site=site)
    preview = preview_cleaner_workbook(uploaded_file)
    if preview["rejected_rows"]:
        return {**preview, "committed": False, "message": "Import stopped. Correct all rejected rows and upload again."}
    created: list[int] = []
    matched: list[int] = []
    assignment_results: list[dict[str, Any]] = []
    with transaction.atomic():
        for item in preview["accepted"]:
            data = dict(item["data"])
            row_number = item["row"]
            identity = {"id_type": data.pop("id_type"), "id_number": data.pop("id_number").upper()}
            cleaner, was_created = Cleaner.objects.get_or_create(
                **identity,
                defaults={**data, "status": status, "created_by": actor, "updated_by": actor},
            )
            (created if was_created else matched).append(cleaner.pk)
            active_assignment = CleanerSiteAssignment.objects.filter(
                cleaner=cleaner, site=site, status__in=["active", "draft"]
            ).exists()
            if not active_assignment:
                assignment = assign_cleaner_to_site(
                    cleaner=cleaner, site=site, assignment_type=CleanerAssignmentType.FULL_TIME,
                    start_date=data["registration_date"], notes=f"Bulk HR {status.value} onboarding row {row_number}.", actor=actor,
                )
                assignment_results.append({"row": row_number, "cleaner_id": cleaner.pk, "site_id": site.pk, "assignment_id": assignment.pk, "status": assignment.status})
            if status == CleanerStatus.TRAINEE and not cleaner.trainee_programs.filter(status__in=["in_training", "extended"]).exists():
                start_trainee_program(
                    cleaner=cleaner, site=site, expected_end_date=data["registration_date"] + timedelta(days=90),
                    start_date=data["registration_date"], actor=actor, notes=f"Created from HR workbook row {row_number}.",
                )
        record_audit(
            action=AuditLog.Action.CREATE, actor=actor, entity=actor,
            summary=f"Imported HR cleaner workbook {preview['file_hash']} to {site.code} as {status.value}",
            after_data={"file_hash": preview["file_hash"], "onboarding_status": status.value, "site_id": site.pk, "created": created, "matched": matched, "assignments": assignment_results},
        )
    return {**preview, "committed": True, "created_cleaner_ids": created, "matched_cleaner_ids": matched, "assignment_results": assignment_results, "onboarding_status": status.value, "site_id": site.pk, "message": f"HR workbook imported successfully to {site.name} as {status.value} cleaners."}


def _assignment_type_for_site(site: Site, current: CleanerSiteAssignment | None) -> CleanerAssignmentType:
    if site.work_mode == WorkMode.SHIFT:
        return CleanerAssignmentType.SHIFT
    if site.work_mode == WorkMode.FULL_TIME:
        return CleanerAssignmentType.FULL_TIME
    return CleanerAssignmentType(current.assignment_type) if current else CleanerAssignmentType.FULL_TIME


def transfer_cleaners_between_sites(
    *,
    cleaner_ids: list[int],
    destination_site: Site,
    effective_date: date,
    reason: str,
    actor: User,
) -> list[dict[str, Any]]:
    """Transfer selected cleaners or trainees as one atomic HR People Registry action.

    Active and draft assignments are ended with history retained. Active trainee
    programmes move to the same destination site so that the receiving site's
    training worksheet immediately becomes the authoritative workspace.
    """
    unique_ids = list(dict.fromkeys(cleaner_ids))
    if not unique_ids:
        raise ValidationError("Select at least one cleaner or trainee to transfer.", code="no_cleaners")
    if len(unique_ids) > 100:
        raise ValidationError("Transfer no more than 100 people at once.", code="batch_too_large")
    if not destination_site.is_active:
        raise ValidationError("Choose an active destination site.", code="destination_inactive")
    if not reason.strip():
        raise ValidationError("A transfer reason is required.", code="reason_required")

    with transaction.atomic():
        cleaners = list(Cleaner.objects.filter(pk__in=unique_ids).order_by("pk"))
        if len(cleaners) != len(unique_ids):
            found_ids = {cleaner.pk for cleaner in cleaners}
            missing = sorted(set(unique_ids) - found_ids)
            raise ValidationError(f"Cleaner records not found: {missing}", code="cleaner_not_found")

        plans: list[tuple[Cleaner, list[CleanerSiteAssignment], TraineeProgram | None]] = []
        active_statuses = [CleanerAssignmentStatus.ACTIVE, CleanerAssignmentStatus.DRAFT]
        trainee_statuses = [TraineeProgramStatus.IN_TRAINING, TraineeProgramStatus.EXTENDED]
        for cleaner in cleaners:
            if cleaner.status == CleanerStatus.INACTIVE:
                raise ValidationError(f"{cleaner.full_name} is inactive and cannot be transferred.", code="cleaner_inactive")
            assignments = list(
                CleanerSiteAssignment.objects.filter(cleaner=cleaner, status__in=active_statuses)
                .select_related("site")
                .order_by("-start_date", "-pk")
            )
            if assignments and all(assignment.site_id == destination_site.pk for assignment in assignments):
                raise ValidationError(
                    f"{cleaner.full_name} is already assigned to {destination_site.name}.", code="already_assigned"
                )
            program = (
                TraineeProgram.objects.filter(cleaner=cleaner, status__in=trainee_statuses)
                .select_related("site")
                .order_by("-start_date", "-pk")
                .first()
            )
            plans.append((cleaner, assignments, program))

        results: list[dict[str, Any]] = []
        for cleaner, assignments, program in plans:
            previous_site = assignments[0].site if assignments else (program.site if program else None)
            for assignment in assignments:
                end_assignment(assignment=assignment, actor=actor, end_date=effective_date)
            assignment = assign_cleaner_to_site(
                cleaner=cleaner,
                site=destination_site,
                assignment_type=_assignment_type_for_site(destination_site, assignments[0] if assignments else None),
                start_date=effective_date,
                notes=f"HR People Registry transfer on {effective_date}: {reason.strip()}",
                actor=actor,
            )
            trainee_program_id: int | None = None
            if program is not None:
                trainee_program_id = transfer_trainee_program(
                    program=program,
                    destination_site=destination_site,
                    effective_date=effective_date,
                    reason=reason,
                    actor=actor,
                ).pk
            elif cleaner.status == CleanerStatus.TRAINEE:
                trainee_program_id = start_trainee_program(
                    cleaner=cleaner,
                    site=destination_site,
                    start_date=effective_date,
                    expected_end_date=effective_date + timedelta(days=90),
                    actor=actor,
                    notes=f"Training programme restored during HR transfer on {effective_date}: {reason.strip()}",
                ).pk
            results.append(
                {
                    "cleaner_id": cleaner.pk,
                    "cleaner_name": cleaner.full_name,
                    "previous_site_id": previous_site.pk if previous_site else None,
                    "previous_site_name": previous_site.name if previous_site else None,
                    "destination_site_id": destination_site.pk,
                    "destination_site_name": destination_site.name,
                    "assignment_id": assignment.pk,
                    "trainee_program_id": trainee_program_id,
                }
            )
        record_audit(
            action=AuditLog.Action.UPDATE,
            actor=actor,
            entity=actor,
            summary=f"Transferred {len(results)} cleaner(s) to {destination_site.name} from HR People Registry",
            after_data={"destination_site_id": destination_site.pk, "effective_date": effective_date.isoformat(), "reason": reason.strip(), "people": results},
        )
    return results
